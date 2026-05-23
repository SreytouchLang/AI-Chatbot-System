import hashlib
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import urlparse

import requests
from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from core.config import get_settings
from core.exceptions import IngestionError
from db.vector import add_documents, file_hash_exists
from utils.transcription_adapter import transcribe_media

settings = get_settings()
SUPPORTED_DOCUMENT_TYPES = {".pdf", ".docx"}
DIRECT_AUDIO_TYPES = {".flac", ".m4a", ".mp3", ".mpeg", ".mpga", ".ogg", ".wav"}
CONVERTIBLE_AUDIO_TYPES = {".aac", ".aif", ".aiff", ".caf", ".wma"}
VIDEO_TYPES = {".avi", ".m4v", ".mkv", ".mov", ".mp4", ".webm"}
SUPPORTED_MEDIA_TYPES = DIRECT_AUDIO_TYPES | CONVERTIBLE_AUDIO_TYPES | VIDEO_TYPES
SUPPORTED_FILE_TYPES = SUPPORTED_DOCUMENT_TYPES | SUPPORTED_MEDIA_TYPES


def clean_text(text: str) -> str:
    text = re.sub(r"\n+", "\n", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _safe_file_name(file_name: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9._-]+", "_", file_name).strip("._")
    return cleaned or "document"


def _guess_suffix(name_or_url: str) -> str:
    if "://" in name_or_url:
        return Path(urlparse(name_or_url).path).suffix.lower()
    return Path(name_or_url).suffix.lower()


def _ensure_supported_suffix(name_or_url: str) -> str:
    suffix = _guess_suffix(name_or_url)
    if suffix not in SUPPORTED_FILE_TYPES:
        supported = ", ".join(sorted(SUPPORTED_FILE_TYPES))
        raise IngestionError(f"Unsupported file type. Supported types: {supported}.")
    return suffix


def _build_display_name(file_name: str, suffix: str) -> str:
    stem = Path(file_name).stem.strip() or "document"
    return f"{stem}{suffix}"


def download_file(source_url: str, file_name: str) -> str:
    suffix = _ensure_supported_suffix(source_url)
    temp_file = tempfile.NamedTemporaryFile(
        prefix=f"{_safe_file_name(file_name)}_",
        suffix=suffix,
        delete=False,
    )

    try:
        response = requests.get(
            source_url,
            stream=True,
            timeout=settings.download_timeout_seconds,
        )
        response.raise_for_status()

        for chunk in response.iter_content(chunk_size=8192):
            if chunk:
                temp_file.write(chunk)
    except requests.RequestException as exc:
        if os.path.exists(temp_file.name):
            os.remove(temp_file.name)
        raise IngestionError(
            "Failed to download the file.",
            details={"source_url": source_url, "reason": str(exc)},
        ) from exc
    finally:
        temp_file.close()
        if "response" in locals():
            response.close()

    return temp_file.name


def create_temp_upload(file_name: str, file_bytes: bytes, suffix: str) -> str:
    temp_file = tempfile.NamedTemporaryFile(
        prefix=f"{_safe_file_name(file_name)}_",
        suffix=suffix,
        delete=False,
    )
    try:
        temp_file.write(file_bytes)
    finally:
        temp_file.close()
    return temp_file.name


def file_hash(file_path: str) -> str:
    sha256 = hashlib.sha256()

    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha256.update(chunk)

    return sha256.hexdigest()


def _ffmpeg_error_message(stderr: str) -> str:
    cleaned = clean_text(stderr)
    if not cleaned:
        return "ffmpeg failed while extracting audio."
    return cleaned[-300:]


def prepare_media_for_transcription(file_path: str) -> tuple[str, bool]:
    suffix = Path(file_path).suffix.lower()
    if suffix in DIRECT_AUDIO_TYPES:
        return file_path, False

    if suffix not in SUPPORTED_MEDIA_TYPES:
        supported = ", ".join(sorted(SUPPORTED_FILE_TYPES))
        raise IngestionError(f"Unsupported file type. Supported types: {supported}.")

    ffmpeg_path = shutil.which("ffmpeg")
    if not ffmpeg_path:
        raise IngestionError(
            "ffmpeg is required to process this media format.",
            details={"file_type": suffix},
        )

    temp_audio = tempfile.NamedTemporaryFile(
        prefix=f"{Path(file_path).stem}_audio_",
        suffix=".mp3",
        delete=False,
    )
    temp_audio.close()

    result = subprocess.run(
        [
            ffmpeg_path,
            "-y",
            "-i",
            file_path,
            "-vn",
            "-ac",
            "1",
            "-ar",
            "16000",
            "-b:a",
            "64k",
            temp_audio.name,
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode == 0:
        return temp_audio.name, True

    if os.path.exists(temp_audio.name):
        os.remove(temp_audio.name)

    raise IngestionError(
        "Failed to extract audio from the media file.",
        details={
            "file_name": Path(file_path).name,
            "reason": _ffmpeg_error_message(result.stderr),
        },
    )


def load_pdf_documents(file_path: str, source_file: str, current_file_hash: str) -> tuple[list[Document], int]:
    pages = PyPDFLoader(file_path).load()
    if not pages:
        raise IngestionError("The PDF was found but no readable pages were extracted.")

    cleaned_pages: list[Document] = []
    for page in pages:
        content = clean_text(page.page_content)
        if not content:
            continue

        metadata = dict(page.metadata or {})
        metadata["source_file"] = source_file
        metadata["file_hash"] = current_file_hash
        metadata["file_type"] = "pdf"
        if isinstance(metadata.get("page"), int):
            metadata["page_number"] = metadata["page"] + 1

        cleaned_pages.append(Document(page_content=content, metadata=metadata))

    if not cleaned_pages:
        raise IngestionError("The PDF does not contain extractable text.")

    return cleaned_pages, len(pages)


def load_docx_documents(file_path: str, source_file: str, current_file_hash: str) -> tuple[list[Document], int]:
    docx_document = DocxDocument(file_path)
    sections: list[Document] = []
    block_index = 0

    for paragraph in docx_document.paragraphs:
        content = clean_text(paragraph.text)
        if not content:
            continue

        block_index += 1
        sections.append(
            Document(
                page_content=content,
                metadata={
                    "source_file": source_file,
                    "file_hash": current_file_hash,
                    "file_type": "docx",
                    "block_index": block_index,
                    "block_type": "paragraph",
                },
            )
        )

    for table_index, table in enumerate(docx_document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if cells:
                rows.append(" | ".join(cells))

        table_text = clean_text("\n".join(rows))
        if not table_text:
            continue

        block_index += 1
        sections.append(
            Document(
                page_content=table_text,
                metadata={
                    "source_file": source_file,
                    "file_hash": current_file_hash,
                    "file_type": "docx",
                    "block_index": block_index,
                    "block_type": "table",
                    "table_index": table_index,
                },
            )
        )

    if not sections:
        raise IngestionError("The DOCX file does not contain extractable text.")

    return sections, 1


def load_media_documents(file_path: str, source_file: str, current_file_hash: str) -> tuple[list[Document], int]:
    suffix = Path(file_path).suffix.lower()
    media_kind = "video" if suffix in VIDEO_TYPES else "audio"
    transcription_input_path, cleanup_required = prepare_media_for_transcription(file_path)

    try:
        transcript = clean_text(transcribe_media(transcription_input_path))
    finally:
        if cleanup_required and os.path.exists(transcription_input_path):
            os.remove(transcription_input_path)

    if not transcript:
        raise IngestionError("The media file did not produce extractable transcript text.")

    return (
        [
            Document(
                page_content=transcript,
                metadata={
                    "source_file": source_file,
                    "file_hash": current_file_hash,
                    "file_type": suffix.lstrip("."),
                    "content_type": "media",
                    "media_kind": media_kind,
                    "block_type": "transcript",
                },
            )
        ],
        1,
    )


def load_documents(file_path: str, source_file: str, current_file_hash: str) -> tuple[list[Document], int, str]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        documents, num_pages = load_pdf_documents(file_path, source_file, current_file_hash)
        return documents, num_pages, "pdf"

    if suffix == ".docx":
        documents, num_pages = load_docx_documents(file_path, source_file, current_file_hash)
        return documents, num_pages, "docx"

    if suffix in SUPPORTED_MEDIA_TYPES:
        documents, num_pages = load_media_documents(file_path, source_file, current_file_hash)
        return documents, num_pages, suffix.lstrip(".")

    supported = ", ".join(sorted(SUPPORTED_FILE_TYPES))
    raise IngestionError(f"Unsupported file type. Supported types: {supported}.")


def store_documents(
    file_name: str,
    file_path: str,
    *,
    source_url: str | None = None,
    input_type: str,
) -> dict[str, object]:
    suffix = _ensure_supported_suffix(file_path)
    display_name = _build_display_name(file_name, suffix)
    current_file_hash = file_hash(file_path)

    if file_hash_exists(current_file_hash):
        return {
            "file_name": display_name,
            "source_url": source_url,
            "status": "skipped",
            "reason": "duplicate content",
            "file_hash": current_file_hash,
            "num_pages": 0,
            "num_chunks": 0,
            "file_type": suffix.lstrip("."),
            "input_type": input_type,
        }

    documents, num_pages, file_type = load_documents(file_path, display_name, current_file_hash)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.ingest_chunk_size,
        chunk_overlap=settings.ingest_chunk_overlap,
    )
    chunks = splitter.split_documents(documents)
    if not chunks:
        raise IngestionError("The file content could not be chunked for ingestion.")

    for index, chunk in enumerate(chunks, start=1):
        chunk.metadata["chunk_index"] = index

    add_documents(chunks)

    return {
        "file_name": display_name,
        "source_url": source_url,
        "num_pages": num_pages,
        "num_chunks": len(chunks),
        "status": "ingested",
        "file_hash": current_file_hash,
        "file_type": file_type,
        "input_type": input_type,
    }


def process_url_file(file_name: str, source_url: str) -> dict[str, object]:
    file_path = download_file(source_url, file_name)

    try:
        return store_documents(
            file_name,
            file_path,
            source_url=source_url,
            input_type="url",
        )
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)


def process_uploaded_file(
    original_filename: str,
    file_bytes: bytes,
    *,
    file_name: str | None = None,
) -> dict[str, object]:
    if not file_bytes:
        raise IngestionError("The uploaded file is empty.")

    suffix = _ensure_supported_suffix(original_filename)
    chosen_name = file_name.strip() if file_name else Path(original_filename).stem
    temp_path = create_temp_upload(chosen_name or "document", file_bytes, suffix)

    try:
        return store_documents(
            chosen_name or "document",
            temp_path,
            input_type="upload",
        )
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
